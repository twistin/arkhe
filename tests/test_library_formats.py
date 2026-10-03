"""DOCX sintético y OCR simulado: procedencia, citas y originales inmutables."""

from io import BytesIO
import json
from pathlib import Path
import subprocess

from docx import Document
from docx.oxml import OxmlElement
import pytest

from docente_ai.doctor import diagnose
from docente_ai.generation.render import render_sources
from docente_ai.generation.validation import validate_response, QuoteResolutionError
from docente_ai.library import ocr
from docente_ai.library.parsers import extract
from docente_ai.library.service import import_document, authorize_document, show_document, delete_document
from docente_ai.rag.chunking import split_segment
from docente_ai.rag.service import citation, index_library, search
from docente_ai.rag.settings import RagSettings
from test_library import library, pdf_bytes
from test_web import web
from test_generation import FakeEmbedder


def docx_bytes():
    document = Document()
    document.add_heading('Programación de música', level=1)
    document.add_paragraph('Contenido teórico ficticio para la sesión.')
    document.add_paragraph('Primer objetivo', style='List Bullet')
    document.add_paragraph('Primera actividad', style='List Number')
    document.add_paragraph('')
    table = document.add_table(rows=2, cols=2)
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row, values in zip(table.rows, [('Tema', 'Criterio'), ('Ritmo | pulso', 'Reconocer el pulso')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph('Conclusión después de la tabla.')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_docx_encabezados_listas_tablas_y_orden():
    result = extract(docx_bytes(), '.docx')
    assert result.status == 'ready', result.error
    texts = [segment['text'] for segment in result.segments]
    assert texts[0] == '# Programación de música'
    assert texts[2] == '- Primer objetivo'
    assert texts[3] == '1. Primera actividad'
    assert texts[4] == '| Tema | Criterio |\n| --- | --- |\n| Ritmo \\| pulso | Reconocer el pulso |'
    assert texts[5] == 'Conclusión después de la tabla.'
    assert result.segments[5]['locator']['paragraph_index'] == 6
    assert result.segments[4]['locator']['table_index'] == 1
    assert all(s['locator']['section'] == 'Programación de música' for s in result.segments)


def test_docx_sin_texto_o_danado_no_se_autoriza():
    stream = BytesIO()
    Document().save(stream)
    assert extract(stream.getvalue(), '.docx').status == 'failed'
    assert extract(b'no es un docx', '.docx').status == 'failed'


def test_docx_importacion_recuperacion_localizadores_y_citas(library):
    db, original = library
    source = original.with_suffix('.docx')
    source.write_bytes(docx_bytes())
    imported = import_document(db, source, category='profesor', subjects=['historia-i'])
    doc = show_document(db, imported['document_id'], text=True)
    assert (db.parent / doc['versions'][0]['original_path']).read_bytes() == source.read_bytes()
    authorize_document(db, doc['id'], enabled=True)
    settings = RagSettings(model='embed:1')
    embedder = FakeEmbedder()
    index_library(db, settings, embedder)
    results = search(db, settings, embedder, 'ritmo', subject='historia-i', category='profesor')['results']
    assert results
    assert all('DOCX' in item['citation'] for item in results)
    for segment in doc['segments']:
        for chunk in split_segment(segment, settings):
            loc = chunk['locator']
            assert chunk['text'] == segment['text'][loc['char_start']:loc['char_end']]
            assert 'Programación de música' in citation(doc['metadata'], loc)


@pytest.fixture
def fake_ocr(monkeypatch):
    calls = []
    output = pdf_bytes('Texto OCR ficticio sobre el ritmo y el pulso musical.')
    monkeypatch.setattr(ocr, 'find_ocr', lambda: '/simulado/ocrmypdf')
    monkeypatch.delenv('DOCENTE_AI_OCR_LANGUAGES', raising=False)

    def run(command, **kwargs):
        assert command[0] == '/simulado/ocrmypdf'
        assert '--force-ocr' in command and '--pages' in command
        assert command[command.index('-l') + 1] == 'spa+eng'
        assert kwargs['timeout'] == 900
        assert kwargs['stderr'] == subprocess.DEVNULL
        calls.append((command, Path(command[-2]).read_bytes()))
        Path(command[-1]).write_bytes(output)
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(ocr.subprocess, 'run', run)
    return calls, output


def test_pdf_sin_ocr_se_rechaza_con_instalacion_y_no_se_autoriza(library):
    db, source = library
    source = source.with_suffix('.pdf')
    source.write_bytes(pdf_bytes(None))
    imported = import_document(db, source, category='profesor', subjects=['historia-i'])
    assert imported['status'] == 'failed'
    assert 'brew install ocrmypdf' in imported['error']
    assert 'doctor --offline' in imported['error']
    with pytest.raises(ValueError, match='versión extraída válida'):
        authorize_document(db, imported['document_id'], enabled=True, accept_warnings=True)


def test_pdf_ocr_derivada_original_revision_y_recuperacion(library, fake_ocr):
    calls, output = fake_ocr
    db, source = library
    source = source.with_suffix('.pdf')
    original = pdf_bytes('12')  # Capa muy escasa: también necesita OCR.
    source.write_bytes(original)
    imported = import_document(db, source, category='documental', subjects=['historia-i'])
    assert imported['status'] == 'needs_review'
    doc = show_document(db, imported['document_id'], text=True)
    native, derived = doc['versions']
    assert derived['id'] == doc['current_version_id']
    assert derived['derived_from_version_id'] == native['id']
    assert derived['text_origin'] == 'texto OCR'
    assert native['text_origin'] == 'texto original'
    assert (db.parent / native['original_path']).read_bytes() == original
    assert (db.parent / derived['original_path']).read_bytes() == output
    assert source.read_bytes() == original
    assert calls[0][1] == original
    assert calls[0][0][calls[0][0].index('--pages') + 1] == '1'
    assert not Path(calls[0][0][-1]).exists()  # Temporales eliminados.
    with pytest.raises(ValueError, match='avisos'):
        authorize_document(db, doc['id'], enabled=True)
    authorize_document(db, doc['id'], enabled=True, accept_warnings=True)
    duplicate = import_document(db, source, category='documental', subjects=['historia-i'])
    assert duplicate['changed'] is False and duplicate['version_id'] == derived['id']
    assert show_document(db, doc['id'])['enabled']
    assert len(show_document(db, doc['id'])['versions']) == 2
    settings = RagSettings(model='embed:1')
    embedder = FakeEmbedder()
    index_library(db, settings, embedder)
    result = search(db, settings, embedder, 'ritmo', subject='historia-i')['results'][0]
    assert result['locator']['text_origin'] == 'texto OCR'
    assert 'posible error de reconocimiento' in result['citation']
    delete_document(db, doc['id'])  # La procedencia no impide eliminar la fuente.


def test_instalar_ocr_permite_recuperar_importacion_fallida(library, monkeypatch):
    db, source = library
    source = source.with_suffix('.pdf')
    source.write_bytes(pdf_bytes(None))
    failed = import_document(db, source, category='profesor', subjects=['historia-i'])
    monkeypatch.setattr(ocr, 'recognize', lambda content, pages: pdf_bytes('Ahora hay texto OCR ficticio suficiente.'))
    recovered = import_document(db, source, category='profesor', subjects=['historia-i'])
    assert recovered['document_id'] == failed['document_id']
    assert recovered['status'] == 'needs_review'
    doc = show_document(db, recovered['document_id'])
    assert len(doc['versions']) == 2
    assert doc['versions'][1]['derived_from_version_id'] == failed['version_id']


def test_original_ocr_alterado_bloquea_autorizacion_y_recuperacion(library, fake_ocr):
    db, source = library
    source = source.with_suffix('.pdf')
    source.write_bytes(pdf_bytes(None))
    imported = import_document(db, source, category='documental', subjects=['historia-i'])
    doc = show_document(db, imported['document_id'])
    authorize_document(db, doc['id'], enabled=True, accept_warnings=True)
    settings = RagSettings(model='embed:1')
    embedder = FakeEmbedder()
    index_library(db, settings, embedder)
    original = db.parent / doc['versions'][0]['original_path']
    original.chmod(0o644)
    original.write_bytes(b'alterado')
    with pytest.raises(ValueError, match='previo al OCR'):
        authorize_document(db, doc['id'], enabled=True, accept_warnings=True)
    with pytest.raises(ValueError, match='previo al OCR'):
        search(db, settings, embedder, 'ritmo', subject='historia-i')


def test_pdf_mixto_ocr_solo_paginas_escasas_y_no_pierde_localizadores(monkeypatch):
    pages_seen = []
    native = 'Contenido nativo suficientemente largo y verificable.'
    def recognize(content, pages):
        pages_seen.extend(pages)
        return pdf_bytes(native, 'Contenido reconocido en la segunda pagina ficticia.')
    monkeypatch.setattr(ocr, 'recognize', recognize)
    result = extract(pdf_bytes(native, None), '.pdf')
    assert pages_seen == [2]
    assert result.segments[0]['text'].strip() == native
    assert result.segments[1]['locator']['pdf_page_index'] == 2
    assert result.segments[1]['locator']['text_origin'] == 'texto OCR'
    assert result.segments[0]['locator']['ocr_document']


@pytest.mark.parametrize('failure', ['returncode', 'timeout', 'invalid', 'pages'])
def test_ocr_fallido_no_crea_texto_autorizable(monkeypatch, failure):
    monkeypatch.setattr(ocr, 'find_ocr', lambda: '/simulado/ocrmypdf')
    def run(command, **kwargs):
        if failure == 'timeout':
            raise subprocess.TimeoutExpired(command, 900)
        if failure == 'invalid':
            Path(command[-1]).write_bytes(b'no es un pdf')
        if failure == 'pages':
            Path(command[-1]).write_bytes(pdf_bytes('Texto suficiente reconocido en la primera pagina.', 'Texto suficiente pero se agrego una segunda pagina.'))
        return subprocess.CompletedProcess(command, 1 if failure == 'returncode' else 0)
    monkeypatch.setattr(ocr.subprocess, 'run', run)
    result = extract(pdf_bytes(None), '.pdf')
    assert result.status == 'failed'
    assert result.derived_content is None


def test_ocr_cita_exacta_y_aviso_en_anexo(fake_ocr):
    parsed = extract(pdf_bytes(None), '.pdf')
    segment = parsed.segments[0]
    source = {'source_id': 'src_1', 'text': segment['text'], 'locator': segment['locator'],
              'category': 'documental', 'metadata': {'title': 'PDF ficticio', 'authors': [], 'year': None}}
    quote = 'Texto OCR ficticio sobre el ritmo y el pulso musical.'
    response = {'status': 'answered', 'claims': [{'kind': 'summary', 'text': 'Síntesis ficticia.',
                'evidence': [{'source_id': 'src_1', 'quote': quote}]}]}
    validated = validate_response(json.dumps(response), [source])
    run = {'evidence': [source], 'result': validated}
    assert 'posible error de reconocimiento' in render_sources(run)
    response['claims'][0]['evidence'][0]['quote'] = quote.replace('ritmo', 'sonido')
    with pytest.raises(QuoteResolutionError):
        validate_response(json.dumps(response), [source])


@pytest.mark.parametrize('available', [False, True])
def test_doctor_comprueba_ocr_sin_ejecutarlo(monkeypatch, available):
    monkeypatch.setattr('docente_ai.doctor.find_ocr', lambda: '/simulado/ocrmypdf' if available else None)
    check = next(item for item in diagnose(offline=True) if item.name == 'ocrmypdf')
    assert check.status == ('ok' if available else 'warning')


def test_web_docx_y_descargas_original_derivada(web, fake_ocr):
    client, ws = web
    uploaded = client.post('/api/upload?name=programacion.docx', content=docx_bytes())
    assert uploaded.status_code == 200, uploaded.text
    path = ws.source_path(uploaded.json()['path'])
    imported = import_document(ws.db, path, category='profesor', subjects=['historia-i'])
    response = client.get('/api/documents/' + imported['document_id'])
    assert response.status_code == 200
    assert response.json()['segments'][0]['locator']['kind'] == 'docx_paragraph'
    scanned = ws.knowledge / 'escaneado.pdf'
    original = pdf_bytes(None)
    scanned.write_bytes(original)
    imported = import_document(ws.db, scanned, category='profesor', subjects=['historia-i'])
    endpoint = '/api/documents/' + imported['document_id']
    assert client.get(endpoint + '?download=1').content == original
    derived = client.get(endpoint + '?download=1&derived=1')
    assert derived.content == fake_ocr[1]
    assert 'texto-ocr.pdf' in derived.headers['Content-Disposition']
    assert client.get(endpoint).json()['versions'][-1]['text_origin'] == 'texto OCR'
    script = client.get('/assets/js/componentes/documentos.js').text
    assert 'posible error de reconocimiento' in script and 'docx_paragraph' in script
